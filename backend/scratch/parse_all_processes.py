import os
import openpyxl
import pandas as pd

base_dir = r"C:\Users\ssanz\OneDrive\Documentos\ChainPointAI\SERVING"
ppto_v5 = os.path.join(base_dir, "20260303 Presupuesto Bosque de Agua V5.xlsx")

wb = openpyxl.load_workbook(ppto_v5, data_only=True)
sheet = wb['Presupuesto V5']

items = []
current_chapter = "PRELIMINARES"

for r in range(11, sheet.max_row + 1):
    c_code = str(sheet.cell(row=r, column=2).value or "").strip()
    c_desc = str(sheet.cell(row=r, column=3).value or "").strip()
    c_unit = str(sheet.cell(row=r, column=4).value or "").strip()
    c_cant = sheet.cell(row=r, column=5).value
    c_total = sheet.cell(row=r, column=9).value

    if c_code.endswith("00") and c_desc and not c_unit:
        current_chapter = c_desc
    elif c_desc and c_total is not None and isinstance(c_total, (int, float)) and c_total > 0:
        items.append({
            "code": c_code,
            "chapter": current_chapter,
            "item": c_desc,
            "unit": c_unit,
            "cant": c_cant,
            "total": c_total
        })

df = pd.DataFrame(items)
print(f"Total ítems válidos en Presupuesto V5: {len(df)}")
print("\nDesglose por Capítulo:")
for cap, group in df.groupby("chapter"):
    print(f"\n=========================================")
    print(f" CAPÍTULO: {cap} ({len(group)} ítems - ${group['total'].sum():,.2f} COP)")
    print(f"=========================================")
    for _, row in group.iterrows():
        print(f"  • [{row['code']}] {row['item']} (Und: {row['unit']}, Cant: {row['cant']}) -> ${row['total']:,.2f} COP")
