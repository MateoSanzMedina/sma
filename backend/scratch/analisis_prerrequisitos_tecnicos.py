import os
import openpyxl
import pandas as pd

base_dir = r"C:\Users\ssanz\OneDrive\Documentos\ChainPointAI\SERVING"
ppto_v5 = os.path.join(base_dir, "20260303 Presupuesto Bosque de Agua V5.xlsx")
dist_flujo = os.path.join(base_dir, "FlujoGerenciaPrueba", "Dist. Costos Flujo.xlsx")
prog_excel = os.path.join(base_dir, "FlujoGerenciaPrueba", "Prog en Excel.xlsx")

# 1. Extraer ítems de Presupuesto V5
items_ppto = []
if os.path.exists(ppto_v5):
    wb = openpyxl.load_workbook(ppto_v5, data_only=True)
    if 'Presupuesto V5' in wb.sheetnames:
        sheet = wb['Presupuesto V5']
        current_chapter = "GENERAL"
        for r in range(11, sheet.max_row + 1):
            code = sheet.cell(row=r, column=2).value
            desc = sheet.cell(row=r, column=3).value
            unit = sheet.cell(row=r, column=4).value
            total = sheet.cell(row=r, column=9).value
            
            if code and isinstance(code, str) and code.endswith("00") and desc:
                current_chapter = str(desc).strip()
            elif desc and total and isinstance(total, (int, float)) and total > 0:
                items_ppto.append({
                    "code": str(code) if code else "",
                    "chapter": current_chapter,
                    "item": str(desc).strip(),
                    "unit": str(unit) if unit else "",
                    "total_cop": total
                })

df_ppto = pd.DataFrame(items_ppto)
print(f"Total ítems extraídos del Presupuesto V5: {len(df_ppto)}")
print(f"Capítulos encontrados ({len(df_ppto['chapter'].unique())}):", list(df_ppto['chapter'].unique()))

# 2. Extraer tareas del Cronograma
tasks_prog = []
if os.path.exists(prog_excel):
    df_prog = pd.read_excel(prog_excel)
    print("Columnas Cronograma Excel:", df_prog.columns.tolist())
    # Tomar la columna de nombre de tarea
    task_col = [c for c in df_prog.columns if 'nombre' in str(c).lower() or 'tarea' in str(c).lower() or 'task' in str(c).lower()][0]
    tasks_prog = df_prog[task_col].dropna().astype(str).str.strip().tolist()
    print(f"Total tareas del cronograma: {len(tasks_prog)}")

# Guardar lista de procesos para categorización
df_ppto.to_csv(os.path.join(base_dir, "SMA", "backend", "scratch", "procesos_presupuesto.csv"), index=False)
