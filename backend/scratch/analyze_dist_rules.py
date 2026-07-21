import openpyxl
import pandas as pd

filepath = r"c:\Users\ssanz\OneDrive\Documentos\ChainPointAI\SERVING\FlujoGerenciaPrueba\Dist. Costos Flujo.xlsx"
wb = openpyxl.load_workbook(filepath, data_only=False) # Get formulas

print("==========================================================================")
print(" REGLAS DE ASOCIACIÓN Y FORMULACIÓN EN Dist. Costos Flujo.xlsx")
print("==========================================================================")

ws_ppto = wb["PPTO ET 4 y 5"]
print("\n--- HOJA 'PPTO ET 4 y 5' (Primeros 30 registros con fórmulas) ---")
for r in range(1, 35):
    row_vals = [ws_ppto.cell(r, c).value for c in range(1, 7)]
    if any(row_vals):
        print(f"Fila {r:2d}: {row_vals}")

ws_flujo = wb["% Flujo"]
print("\n--- HOJA '% Flujo' (Secciones y Distribución de Tramos / Precursores) ---")
for r in range(1, 40):
    row_vals = [ws_flujo.cell(r, c).value for c in range(1, 8)]
    if any(row_vals):
        print(f"Fila {r:2d}: {row_vals}")

ws_apu = wb["APU"]
print("\n--- HOJA 'APU' (Composición de Insumos y Requisitos por Ítem) ---")
for r in range(1, 35):
    row_vals = [ws_apu.cell(r, c).value for c in range(1, 10)]
    if any(row_vals):
        print(f"Fila {r:2d}: {row_vals}")
