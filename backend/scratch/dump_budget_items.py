import os
import openpyxl
import json

ppto_v5 = r"C:\Users\ssanz\OneDrive\Documentos\ChainPointAI\SERVING\20260303 Presupuesto Bosque de Agua V5.xlsx"
wb = openpyxl.load_workbook(ppto_v5, data_only=True)
sheet = wb['Presupuesto V5']

chapters = {}
current_chapter = "PRELIMINARES Y ADECUACION"

for r in range(11, sheet.max_row + 1):
    c_code = str(sheet.cell(row=r, column=2).value or "").strip()
    c_desc = str(sheet.cell(row=r, column=3).value or "").strip()
    c_unit = str(sheet.cell(row=r, column=4).value or "").strip()
    c_total = sheet.cell(row=r, column=9).value

    if c_code.endswith("00") and c_desc and not c_unit:
        current_chapter = c_desc
        if current_chapter not in chapters:
            chapters[current_chapter] = []
    elif c_desc and c_total is not None and isinstance(c_total, (int, float)) and c_total > 0:
        if current_chapter not in chapters:
            chapters[current_chapter] = []
        chapters[current_chapter].append({
            "code": c_code,
            "name": c_desc,
            "unit": c_unit,
            "cost": c_total
        })

print(f"Total capítulos encontrados: {len(chapters)}")
for ch, items in chapters.items():
    print(f"  • {ch}: {len(items)} ítems")

out_file = r"C:\Users\ssanz\OneDrive\Documentos\ChainPointAI\SERVING\SMA\backend\scratch\all_items_dump.json"
with open(out_file, "w", encoding="utf-8") as f:
    json.dump(chapters, f, ensure_ascii=False, indent=2)

print("\nDump guardado en:", out_file)
