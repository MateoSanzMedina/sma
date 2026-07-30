import os
import openpyxl
import pandas as pd

ppto_path = r"C:\Users\ssanz\OneDrive\Documentos\ChainPointAI\SERVING\20260303 Presupuesto Bosque de Agua V5.xlsx"
dist_path = r"C:\Users\ssanz\OneDrive\Documentos\ChainPointAI\SERVING\FlujoGerenciaPrueba\Dist. Costos Flujo.xlsx"

print("--- 20260303 Presupuesto Bosque de Agua V5.xlsx ---")
if os.path.exists(ppto_path):
    wb = openpyxl.load_workbook(ppto_path, data_only=True)
    print("Sheets:", wb.sheetnames[:10])
    for sname in wb.sheetnames[:5]:
        sheet = wb[sname]
        print(f"Sheet '{sname}': max_row={sheet.max_row}, max_col={sheet.max_column}")
        for r in range(1, min(10, sheet.max_row + 1)):
            row_vals = [sheet.cell(row=r, column=c).value for c in range(1, min(10, sheet.max_column + 1))]
            print(f"  Row {r}: {row_vals}")

print("\n--- Dist. Costos Flujo.xlsx ---")
if os.path.exists(dist_path):
    wb2 = openpyxl.load_workbook(dist_path, data_only=True)
    print("Sheets:", wb2.sheetnames[:10])
    for sname in wb2.sheetnames[:5]:
        sheet = wb2[sname]
        print(f"Sheet '{sname}': max_row={sheet.max_row}, max_col={sheet.max_column}")
        for r in range(1, min(10, sheet.max_row + 1)):
            row_vals = [sheet.cell(row=r, column=c).value for c in range(1, min(10, sheet.max_column + 1))]
            print(f"  Row {r}: {row_vals}")
