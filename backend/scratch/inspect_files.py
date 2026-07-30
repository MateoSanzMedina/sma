import os
import openpyxl
import pandas as pd

dir_path = r"C:\Users\ssanz\OneDrive\Documentos\ChainPointAI\SERVING\FlujoGerenciaPrueba"
files = [
    "Dist. Costos Flujo.xlsx",
    "informe_flujo_caja_serving_2026-07-21 (3).xlsx",
    "informe_flujo_caja_serving_2026-07-21 (4).xlsx"
]

for f in files:
    full_p = os.path.join(dir_path, f)
    if os.path.exists(full_p):
        print(f"=== FILE: {f} ===")
        wb = openpyxl.load_workbook(full_p, data_only=True)
        print("Sheets:", wb.sheetnames)
        for sname in wb.sheetnames[:3]:
            sheet = wb[sname]
            print(f"  Sheet: {sname}, max_row={sheet.max_row}, max_col={sheet.max_column}")
            # print first 5 rows
            for r in range(1, min(6, sheet.max_row + 1)):
                row_vals = [sheet.cell(row=r, column=c).value for c in range(1, min(10, sheet.max_column + 1))]
                print(f"    Row {r}: {row_vals}")
