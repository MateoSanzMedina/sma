import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
import io
import re
import math
import base64
from copy import copy

# ---------------------------------------------------------------------------
# Constants – canonical header keywords used to locate columns.
# Identical to analisis_apus_final.py
# ---------------------------------------------------------------------------
SHEET0_ITEM_KEYWORDS   = ['ítem', 'item', 'código', 'codigo']   # col with APU codes (sheet 0)
SHEET0_ACCUM_KEYWORDS  = ['cant', 'ejecut', 'acumul']           # col with accumulated quantity (sheet 0, ALL must match)
SHEET1_CODE_KEYWORDS   = ['código', 'codigo']                    # col with APU codes (sheet 1)
SHEET1_INSUMO_KEYWORDS = ['insumo']                              # col with insumo codes (sheet 1)
SHEET1_NAME_KEYWORDS   = ['descripción', 'descripcion', 'nombre', 'concepto']
SHEET1_GROUP_KEYWORDS  = ['análisis unitario ejecución',
                           'analisis unitario ejecucion']         # group header spanning execution section
SHEET1_SUBHDR_KEYWORDS = ['cant', 'ejecut']                      # subheader col label inside execution section (ALL must match)

HEADER_SEARCH_ROWS = 30
TOLERANCIA_CONCRETO = 0.15


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def copy_cell_style(src, dst):
    """Copy font, border, fill, number_format, protection and alignment."""
    if src and src.has_style and type(dst).__name__ != 'MergedCell':
        dst.font          = copy(src.font)
        dst.border        = copy(src.border)
        dst.fill          = copy(src.fill)
        dst.number_format = src.number_format
        dst.protection    = copy(src.protection)
        dst.alignment     = copy(src.alignment)


def normalize(val):
    if val and isinstance(val, str):
        return ' '.join(val.lower().split())
    return ''


def all_keywords_in(text, keywords):
    return all(k in text for k in keywords)


def find_header_cell(ws, keywords, match_all=True, max_row=HEADER_SEARCH_ROWS, max_col=None):
    max_col = max_col or ws.max_column
    for r in range(1, min(max_row + 1, ws.max_row + 1)):
        for c in range(1, max_col + 1):
            text = normalize(ws.cell(r, c).value)
            if not text:
                continue
            hit = all_keywords_in(text, keywords) if match_all else any(k in text for k in keywords)
            if hit:
                return r, c
    return None, None


def find_column_in_row(ws, row, keywords, match_all=True, col_start=1, col_end=None):
    col_end = col_end or ws.max_column
    for c in range(col_start, col_end + 1):
        text = normalize(ws.cell(row, c).value)
        if not text:
            continue
        hit = all_keywords_in(text, keywords) if match_all else any(k in text for k in keywords)
        if hit:
            return c
    return None


def clean_numeric(val):
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        if math.isnan(val):
            return 0.0
        return float(val)
    val_str = str(val).replace("$", "").replace(",", "").replace(" ", "").strip()
    try:
        return float(val_str)
    except ValueError:
        return 0.0


# ---------------------------------------------------------------------------
# Sheet-0 reader (Hoja 0)
# ---------------------------------------------------------------------------
def read_accumulated_values(wb):
    ws = wb.worksheets[0]
    item_row, item_col = find_header_cell(ws, SHEET0_ITEM_KEYWORDS, match_all=False)
    accum_row, accum_col = find_header_cell(ws, SHEET0_ACCUM_KEYWORDS, match_all=True)

    if item_col is None or accum_col is None:
        item_row, item_col = 7, 2
        accum_row, accum_col = 7, 10

    data_start = max(item_row or 0, accum_row or 0) + 1
    accum_dict = {}

    for r in range(data_start, ws.max_row + 1):
        code_val  = ws.cell(r, item_col).value
        accum_val = ws.cell(r, accum_col).value
        if code_val is None:
            continue
        code_str = str(code_val).strip()
        if not code_str:
            continue
        try:
            accum_dict[code_str] = float(accum_val) if accum_val is not None else 0.0
        except (ValueError, TypeError):
            accum_dict[code_str] = clean_numeric(accum_val)

    return accum_dict


# ---------------------------------------------------------------------------
# Sheet-1 Header Detection
# ---------------------------------------------------------------------------
def detect_sheet1_headers(ws):
    code_row, code_col = find_header_cell(ws, SHEET1_CODE_KEYWORDS, match_all=False)
    insumo_row, insumo_col = find_header_cell(ws, SHEET1_INSUMO_KEYWORDS, match_all=False)
    group_row, group_col = find_header_cell(ws, SHEET1_GROUP_KEYWORDS, match_all=False)

    ejecutada_col = None
    subheader_row = None

    if group_row is not None and group_col is not None:
        for r in range(group_row + 1, min(group_row + HEADER_SEARCH_ROWS, ws.max_row + 1)):
            col = find_column_in_row(ws, r, SHEET1_SUBHDR_KEYWORDS, match_all=True, col_start=group_col)
            if col is not None:
                ejecutada_col = col
                subheader_row = r
                break

    if ejecutada_col is None:
        subheader_row, ejecutada_col = find_header_cell(ws, SHEET1_SUBHDR_KEYWORDS, match_all=True)

    unitaria_col = None
    precio_col   = None

    if subheader_row is not None and ejecutada_col is not None:
        for c in range(ejecutada_col + 1, ws.max_column + 1):
            text = normalize(ws.cell(subheader_row, c).value)
            if unitaria_col is None and 'unitaria' in text:
                unitaria_col = c
            if precio_col is None and 'precio' in text:
                precio_col = c
            if unitaria_col and precio_col:
                break

    if code_col is None:
        code_row, code_col = 2, 1
    if insumo_col is None:
        insumo_col = code_col + 1
    if ejecutada_col is None or subheader_row is None:
        subheader_row, ejecutada_col = 3, 9
    if unitaria_col is None:
        unitaria_col = ejecutada_col + 1
    if precio_col is None:
        precio_col = ejecutada_col + 2

    group_header_row = group_row if group_row is not None else subheader_row
    name_col = code_col + 2

    base_group_col = 5
    for c in range(1, ws.max_column + 1):
        v = normalize(ws.cell(group_header_row, c).value)
        if 'base' in v:
            base_group_col = c
            break

    return {
        'code_col':         code_col,
        'insumo_col':       insumo_col,
        'name_col':         name_col,
        'group_header_row': group_header_row,
        'subheader_row':    subheader_row,
        'base_group_col':   base_group_col,
        'ejecutada_col':    ejecutada_col,
        'unitaria_col':     unitaria_col,
        'precio_col':       precio_col,
    }


# ---------------------------------------------------------------------------
# Map accumulated & insert calculated columns in workbook (Paso 1)
# Exactly matching analisis_apus_final.py
# ---------------------------------------------------------------------------
def map_accumulated_wb(wb, accum_dict):
    ws = wb.worksheets[1] if len(wb.worksheets) > 1 else wb.worksheets[0]
    h = detect_sheet1_headers(ws)

    code_col         = h['code_col']
    insumo_col       = h['insumo_col']
    group_header_row = h['group_header_row']
    subheader_row    = h['subheader_row']
    ejecutada_col    = h['ejecutada_col']
    unitaria_col     = h['unitaria_col']
    precio_col       = h['precio_col']

    # Step 3a: find or reuse 'Obra faltante'
    col_accum = None
    for c in range(1, ws.max_column + 1):
        if normalize(ws.cell(group_header_row, c).value) == 'obra faltante':
            col_accum = c
            break
    if col_accum is None:
        last_col = 1
        for c in range(1, ws.max_column + 1):
            v1 = ws.cell(group_header_row, c).value
            v2 = ws.cell(subheader_row, c).value
            if v1 is not None or v2 is not None:
                last_col = c
        col_accum = last_col + 1
    col_accum_letter = get_column_letter(col_accum)

    # Step 3b: find or reuse 'Proporción'
    col_ratio = None
    for c in range(1, ws.max_column + 1):
        if normalize(ws.cell(group_header_row, c).value) == 'proporcion':
            col_ratio = c
            break
    if col_ratio is None:
        col_ratio = col_accum + 1
    col_ratio_letter = get_column_letter(col_ratio)

    # Step 3c: find or reuse 'Cant. Unitaria'
    col_unitaria = None
    for c in range(1, ws.max_column + 1):
        if normalize(ws.cell(subheader_row, c).value) in ('cant. unitaria ejecucion', 'cant. unitaria ejec'):
            col_unitaria = c
            break
    if col_unitaria is None:
        col_unitaria = col_ratio + 1
    col_unitaria_letter = get_column_letter(col_unitaria)

    # Step 3d: find or reuse 'Precio'
    col_precio = None
    for c in range(1, ws.max_column + 1):
        if normalize(ws.cell(subheader_row, c).value) in ('precio ejecucion', 'precio ejec'):
            col_precio = c
            break
    if col_precio is None:
        col_precio = col_unitaria + 1
    col_precio_letter = get_column_letter(col_precio)

    # Step 3e: find or reuse 'Costo Ejec'
    col_costo = None
    for c in range(1, ws.max_column + 1):
        if normalize(ws.cell(subheader_row, c).value) in ('costo ejec', 'costo ejecucion'):
            col_costo = c
            break
    if col_costo is None:
        col_costo = col_precio + 1
    col_costo_letter = get_column_letter(col_costo)

    # Step 3f: find or reuse 'Total Unitaria'
    col_total = None
    for c in range(1, ws.max_column + 1):
        if normalize(ws.cell(subheader_row, c).value) in ('total unitaria', 'total unid'):
            col_total = c
            break
    if col_total is None:
        col_total = col_costo + 1
    col_total_letter = get_column_letter(col_total)

    # Step 3g: find or reuse 'Diferencia'
    col_diff = None
    for c in range(1, ws.max_column + 1):
        if normalize(ws.cell(subheader_row, c).value) in ('diferencia', 'dif'):
            col_diff = c
            break
    if col_diff is None:
        col_diff = col_total + 1
    col_diff_letter = get_column_letter(col_diff)

    # Step 3h: find or reuse 'Estado'
    col_status = None
    for c in range(1, ws.max_column + 1):
        if normalize(ws.cell(subheader_row, c).value) in ('estado', 'estado diferencia', 'observacion', 'resultado'):
            col_status = c
            break
    if col_status is None:
        col_status = col_diff + 1
    col_status_letter = get_column_letter(col_status)

    # Group-level headers
    grp_src = ws.cell(group_header_row, ejecutada_col)
    for col, label in [
        (col_accum,    "Obra faltante"),
        (col_ratio,    "Proporcion"),
        (col_unitaria, "Analisis unitario ejecucion"),
        (col_precio,   "Analisis unitario ejecucion"),
        (col_costo,    "Analisis unitario ejecucion"),
        (col_total,    "Analisis unitario ejecucion"),
        (col_diff,     "Analisis unitario ejecucion"),
        (col_status,   "Analisis unitario ejecucion"),
    ]:
        cell = ws.cell(group_header_row, col)
        cell.value = label
        copy_cell_style(grp_src, cell)

    # Subheaders
    sub_src = ws.cell(subheader_row, ejecutada_col)
    for col, label in [
        (col_accum,    "Cant. Ejecutada"),
        (col_ratio,    "Sub / APU"),
        (col_unitaria, "Cant. Unitaria Ejec"),
        (col_precio,   "Precio Ejec"),
        (col_costo,    "Costo Ejec"),
        (col_total,    "Total Unitaria"),
        (col_diff,     "Diferencia"),
        (col_status,   "Estado"),
    ]:
        cell = ws.cell(subheader_row, col)
        cell.value = label
        copy_cell_style(sub_src, cell)

    data_start_row = subheader_row + 1
    current_parent_row = None

    for r in range(data_start_row, ws.max_row + 1):
        code   = ws.cell(r, code_col).value
        insumo = ws.cell(r, insumo_col).value

        code_str = str(code).strip() if code is not None else ""
        if not code_str:
            continue

        is_parent = insumo is None or str(insumo).strip() == "" or str(insumo).lower() in ("none", "nan")
        src_cell  = ws.cell(r, ejecutada_col)

        if is_parent:
            raw_ejec = src_cell.value
            is_empty_ejec = (raw_ejec is None or str(raw_ejec).strip() in ('', '0', '0.0') or raw_ejec == 0)
            if is_empty_ejec and code_str in accum_dict and accum_dict[code_str] is not None:
                src_cell.value = accum_dict[code_str]

            if code_str in accum_dict:
                dst_cell = ws.cell(r, col_accum)
                copy_cell_style(src_cell, dst_cell)
                dst_cell.value = accum_dict[code_str]
            current_parent_row = r

            for col in (col_ratio, col_unitaria, col_precio, col_costo, col_total, col_diff, col_status):
                ws.cell(r, col).value = None
        else:
            raw_val = src_cell.value
            dst_cell = ws.cell(r, col_accum)
            if raw_val is not None:
                copy_cell_style(src_cell, dst_cell)
                dst_cell.value = raw_val

            ejec_letter = get_column_letter(ejecutada_col)
            u_src = ws.cell(r, unitaria_col)
            if current_parent_row is not None:
                u_src.value = (
                    f"=IF({ejec_letter}{current_parent_row}<>0, "
                    f"{ejec_letter}{r}/{ejec_letter}{current_parent_row}, 0)"
                )

            ratio_cell = ws.cell(r, col_ratio)
            copy_cell_style(src_cell, ratio_cell)
            if current_parent_row is not None:
                ratio_cell.value = (
                    f"=IF({col_accum_letter}{current_parent_row}<>0,"
                    f"{col_accum_letter}{r}/{col_accum_letter}{current_parent_row},\"\")"
                )
            else:
                ratio_cell.value = None

            u_cell = ws.cell(r, col_unitaria)
            copy_cell_style(u_src, u_cell)
            u_cell.value = u_src.value

            p_src  = ws.cell(r, precio_col)
            p_cell = ws.cell(r, col_precio)
            copy_cell_style(p_src, p_cell)
            p_cell.value = p_src.value

            c_cell = ws.cell(r, col_costo)
            copy_cell_style(p_src, c_cell)
            if current_parent_row is not None:
                c_cell.value = (
                    f'=IF({col_ratio_letter}{r}<>"", {col_precio_letter}{r}*{col_ratio_letter}{r}, "")'
                )

            t_cell = ws.cell(r, col_total)
            copy_cell_style(u_src, t_cell)
            if current_parent_row is not None:
                t_cell.value = (
                    f'=IF({col_accum_letter}{current_parent_row}<>0, {col_unitaria_letter}{r}*{col_accum_letter}{current_parent_row}, "")'
                )

            d_cell = ws.cell(r, col_diff)
            copy_cell_style(u_src, d_cell)
            if current_parent_row is not None:
                d_cell.value = (
                    f'=IF({col_accum_letter}{current_parent_row}<>0, {col_total_letter}{r}-{col_accum_letter}{r}, "")'
                )

            st_cell = ws.cell(r, col_status)
            copy_cell_style(u_src, st_cell)
            if current_parent_row is not None:
                st_cell.value = (
                    f'=IF(ISNUMBER({col_diff_letter}{r}), '
                    f'IF({col_diff_letter}{r}>0, "sobra", '
                    f'IF({col_diff_letter}{r}<0, "falta", "no coincide")), '
                    f'"no coincide")'
                )

    ws.column_dimensions[col_accum_letter].width    = 20
    ws.column_dimensions[col_ratio_letter].width    = 15
    ws.column_dimensions[col_unitaria_letter].width = 20
    ws.column_dimensions[col_precio_letter].width   = 18
    ws.column_dimensions[col_costo_letter].width    = 18
    ws.column_dimensions[col_total_letter].width    = 18
    ws.column_dimensions[col_diff_letter].width     = 18
    ws.column_dimensions[col_status_letter].width   = 18
    return True


# ---------------------------------------------------------------------------
# Concrete Validation Sheet (Paso 2)
# Exactly matching analisis_apus_final.py
# ---------------------------------------------------------------------------
def find_mix_design_sheets(wb, exclude_titles):
    fichas = []
    for ws in wb.worksheets:
        if ws.title in exclude_titles:
            continue

        insumo_row, insumo_col = find_header_cell(ws, ['insumo'], match_all=False)
        if insumo_col is None:
            continue
        qty_row, qty_col = find_header_cell(ws, ['cantidad'], match_all=False)
        if qty_col is None:
            qty_col = insumo_col + 2

        rows = {}
        for r in range(insumo_row + 1, ws.max_row + 1):
            name = normalize(ws.cell(r, insumo_col).value)
            if not name:
                continue
            if 'cemento' in name and 'cemento' not in rows:
                rows['cemento'] = r
            elif 'arena' in name and 'arena' not in rows:
                rows['arena'] = r
            elif 'triturado' in name and 'triturado' not in rows:
                rows['triturado'] = r

        if {'cemento', 'arena', 'triturado'} <= rows.keys():
            fichas.append({
                'sheet': ws,
                'title': ws.title,
                'cemento_row': rows['cemento'],
                'arena_row': rows['arena'],
                'triturado_row': rows['triturado'],
                'qty_col': qty_col,
            })
    return fichas


def find_concrete_apus(ws, code_col, insumo_col, nombre_col, ejec_col_real):
    items = []
    current = None

    for r in range(1, ws.max_row + 1):
        code = ws.cell(r, code_col).value
        insumo = ws.cell(r, insumo_col).value
        nombre = ws.cell(r, nombre_col).value

        code_str = str(code).strip() if code is not None else ""
        if not code_str:
            continue

        is_parent = insumo is None or str(insumo).strip() == "" or str(insumo).lower() in ("none", "nan")

        if is_parent:
            if current and {'cemento', 'arena', 'triturado'} <= current['rows'].keys():
                items.append(current)
            current = {'codigo': code_str, 'nombre_item': str(nombre or ""),
                       'fila_padre': r, 'rows': {}}
            continue

        if current is None:
            continue

        name_norm = normalize(str(nombre))
        if 'subtotales' in name_norm or 'subtotal' in name_norm:
            continue
        if 'cemento' in name_norm and 'cemento' not in current['rows']:
            current['rows']['cemento'] = r
        elif 'arena' in name_norm and 'arena' not in current['rows']:
            current['rows']['arena'] = r
        elif 'triturado' in name_norm and 'triturado' not in current['rows']:
            current['rows']['triturado'] = r

    if current and {'cemento', 'arena', 'triturado'} <= current['rows'].keys():
        items.append(current)

    return items


def col_letter(ws_title, col, row):
    return f"'{ws_title}'!{get_column_letter(col)}{row}"


def build_validation_sheet_wb(wb):
    apu_sheet_title = None
    ws_apu = None
    for ws in wb.worksheets:
        r, c = find_header_cell(ws, ['análisis unitario ejecución', 'analisis unitario ejecucion'], match_all=False)
        if c is not None:
            apu_sheet_title = ws.title
            ws_apu = ws
            break

    if ws_apu is None:
        if len(wb.worksheets) > 1:
            ws_apu = wb.worksheets[1]
            apu_sheet_title = ws_apu.title
        else:
            return

    h = detect_sheet1_headers(ws_apu)
    code_col = h['code_col']
    insumo_col = h['insumo_col']
    nombre_col = code_col + 2
    ejec_col_real = h['ejecutada_col']

    col_obra_faltante = None
    for c in range(1, ws_apu.max_column + 1):
        if normalize(ws_apu.cell(h['group_header_row'], c).value) == 'obra faltante':
            col_obra_faltante = c
            break
    qty_col = col_obra_faltante if col_obra_faltante else ejec_col_real

    fichas = find_mix_design_sheets(wb, exclude_titles={apu_sheet_title, wb.worksheets[0].title})
    if not fichas:
        return

    items = find_concrete_apus(ws_apu, code_col, insumo_col, nombre_col, qty_col)
    if not items:
        return

    sheet_name = "Validacion Mezcla Concreto"
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    vs = wb.create_sheet(sheet_name)

    bold = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style='thin', color='B7B7B7')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical='center', horizontal='center')

    vs['A1'] = "Validación de mezcla de concreto: Cemento vs. Arena y Triturado"
    vs['A1'].font = Font(bold=True, size=13)
    vs['A2'] = ("Compara el cemento REGISTRADO en cada APU contra el cemento que "
                "deberian implicar la arena y el triturado registrados, segun la "
                "ficha de diseno de mezcla. Tolerancia: ±{:.0f}%.").format(TOLERANCIA_CONCRETO * 100)
    vs['A2'].font = Font(italic=True, size=9, color="555555")

    headers = [
        "Código Item", "Nombre Item", "Ficha usada",
        "Cemento registrado (sc)", "Arena registrada (m3)", "Triturado registrado (m3)",
        "Cemento esperado según Arena (sc)", "Cemento esperado según Triturado (sc)",
        "Dif. % Cemento vs. Arena", "Dif. % Cemento vs. Triturado",
        "Dif. % Arena vs Triturado (proporción)", "Estado"
    ]
    header_row = 4
    for j, htext in enumerate(headers, start=1):
        cell = vs.cell(header_row, j, htext)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap
        cell.border = border

    row = header_row + 1
    for item in items:
        ficha = fichas[0]
        item_text = normalize(item['nombre_item'])
        for f in fichas:
            digits = ''.join(ch for ch in f['title'] if ch.isdigit())
            if digits and digits in item_text:
                ficha = f
                break

        f_title = ficha['sheet'].title
        f_qty_col = ficha['qty_col']
        cemento_ficha_ref = col_letter(f_title, f_qty_col, ficha['cemento_row'])
        arena_ficha_ref = col_letter(f_title, f_qty_col, ficha['arena_row'])
        triturado_ficha_ref = col_letter(f_title, f_qty_col, ficha['triturado_row'])

        cemento_real_ref = col_letter(apu_sheet_title, qty_col, item['rows']['cemento'])
        arena_real_ref = col_letter(apu_sheet_title, qty_col, item['rows']['arena'])
        triturado_real_ref = col_letter(apu_sheet_title, qty_col, item['rows']['triturado'])

        vs.cell(row, 1, item['codigo'])
        vs.cell(row, 2, item['nombre_item'])
        vs.cell(row, 3, f_title)
        vs.cell(row, 4, f"={cemento_real_ref}")
        vs.cell(row, 5, f"={arena_real_ref}")
        vs.cell(row, 6, f"={triturado_real_ref}")
        vs.cell(row, 7, f"=IFERROR({arena_real_ref}*({cemento_ficha_ref}/{arena_ficha_ref}),\"\")")
        vs.cell(row, 8, f"=IFERROR({triturado_real_ref}*({cemento_ficha_ref}/{triturado_ficha_ref}),\"\")")
        col_g = get_column_letter(7)
        col_h = get_column_letter(8)
        vs.cell(row, 9, f"=IFERROR(({cemento_real_ref}-{col_g}{row})/{col_g}{row},\"\")")
        vs.cell(row, 9).number_format = '0.0%'
        vs.cell(row, 10, f"=IFERROR(({cemento_real_ref}-{col_h}{row})/{col_h}{row},\"\")")
        vs.cell(row, 10).number_format = '0.0%'
        vs.cell(row, 11,
                f"=IFERROR((({arena_real_ref}/{triturado_real_ref})-({arena_ficha_ref}/{triturado_ficha_ref}))"
                f"/({arena_ficha_ref}/{triturado_ficha_ref}),\"\")")
        vs.cell(row, 11).number_format = '0.0%'
        col_i = get_column_letter(9)
        col_j = get_column_letter(10)
        col_k = get_column_letter(11)
        vs.cell(row, 12,
                f'=IF(OR(ABS({col_i}{row})>{TOLERANCIA_CONCRETO},ABS({col_j}{row})>{TOLERANCIA_CONCRETO},'
                f'ABS({col_k}{row})>{TOLERANCIA_CONCRETO}),"Revisar","OK")')

        for c in range(1, 13):
            vs.cell(row, c).border = border
            if c in (4, 5, 6, 7, 8):
                vs.cell(row, c).number_format = '0.00'

        row += 1

    last_row = row - 1
    if last_row >= header_row + 1:
        vs.conditional_formatting.add(
            f"L{header_row+1}:L{last_row}",
            CellIsRule(operator='equal', formula=['"Revisar"'],
                       fill=PatternFill("solid", fgColor="F8CBAD"))
        )
        vs.conditional_formatting.add(
            f"L{header_row+1}:L{last_row}",
            CellIsRule(operator='equal', formula=['"OK"'],
                       fill=PatternFill("solid", fgColor="C6E0B4"))
        )

    widths = [14, 32, 14, 20, 18, 20, 22, 24, 20, 22, 24, 12]
    for j, w in enumerate(widths, start=1):
        vs.column_dimensions[get_column_letter(j)].width = w
    vs.freeze_panes = "A5"


# ---------------------------------------------------------------------------
# Main processing function for FastAPI Service
# ---------------------------------------------------------------------------
async def process_sao_costs(file_content: bytes, filename: str):
    """
    Procesa el archivo Excel de SAO ejecutando exactamente el mismo flujo
    que analisis_apus_final.py. Retorna JSON + excel_b64.
    """
    try:
        # 1. Leer acumulados con data_only=True de Hoja 0 (idéntico a Paso 1 del script)
        wb_ro = openpyxl.load_workbook(io.BytesIO(file_content), data_only=True)
        accum_dict = read_accumulated_values(wb_ro)
        wb_ro.close()

        # 2. Cargar archivo para edición (idéntico a Paso 1 del script)
        wb = openpyxl.load_workbook(io.BytesIO(file_content))
    except Exception as e:
        raise ValueError(f"No se pudo leer el archivo Excel con openpyxl: {str(e)}")

    if not wb.sheetnames:
        raise ValueError("El archivo Excel no contiene hojas de cálculo.")

    # 3. Ejecutar Paso 1: map_accumulated_in_place
    map_accumulated_wb(wb, accum_dict)

    # 4. Ejecutar Paso 2: build_validation_sheet
    build_validation_sheet_wb(wb)

    # 5. Extraer objetos para la respuesta JSON (APUs, alertas, resumen, etc.)
    ws_apu = wb.worksheets[1] if len(wb.worksheets) > 1 else wb.worksheets[0]
    h = detect_sheet1_headers(ws_apu)

    code_col      = h['code_col']
    insumo_col    = h['insumo_col']
    name_col      = h['name_col']
    subheader_row = h['subheader_row']
    base_group_col = h['base_group_col']
    ejecutada_col = h['ejecutada_col']
    unitaria_col  = h['unitaria_col']
    precio_col    = h['precio_col']

    col_obra_faltante = None
    for c in range(1, ws_apu.max_column + 1):
        if normalize(ws_apu.cell(h['group_header_row'], c).value) == 'obra faltante':
            col_obra_faltante = c
            break

    apus = []
    current_apu = None

    for r in range(subheader_row + 1, ws_apu.max_row + 1):
        code_val = ws_apu.cell(r, code_col).value
        insumo_val = ws_apu.cell(r, insumo_col).value
        name_val = ws_apu.cell(r, name_col).value

        code_str = str(code_val).strip() if code_val is not None else ""
        if not code_str:
            continue

        clean_code = code_str.replace(".0", "").replace(".", "").strip()
        if not re.match(r'^\d+$', clean_code):
            continue

        insumo_str = str(insumo_val).strip() if insumo_val is not None else ""
        nombre = str(name_val).strip() if name_val is not None else ""
        unit_cell = ws_apu.cell(r, name_col + 1).value
        unidad = str(unit_cell).strip() if unit_cell is not None else ""

        is_parent = not insumo_str or insumo_str.lower() in ("none", "nan", "")

        if is_parent:
            base_qty = clean_numeric(ws_apu.cell(r, base_group_col).value)
            ejec_qty = clean_numeric(ws_apu.cell(r, ejecutada_col).value)

            if col_obra_faltante is not None:
                obra_faltante = clean_numeric(ws_apu.cell(r, col_obra_faltante).value)
            elif clean_code in accum_dict:
                obra_faltante = accum_dict[clean_code]
            else:
                obra_faltante = max(0.0, base_qty - ejec_qty) if base_qty > ejec_qty else ejec_qty

            current_apu = {
                "code": clean_code,
                "name": nombre,
                "unit": unidad,
                "base_qty": base_qty,
                "ejec_qty": ejec_qty,
                "faltante_qty_teorica": obra_faltante,
                "obra_faltante_qty": obra_faltante,
                "resources": []
            }
            apus.append(current_apu)

        elif insumo_str != "99999" and "subtotal" not in insumo_str.lower() and current_apu:
            base_qty = clean_numeric(ws_apu.cell(r, base_group_col).value)
            base_unit_qty = clean_numeric(ws_apu.cell(r, base_group_col + 1).value)
            base_price = clean_numeric(ws_apu.cell(r, base_group_col + 2).value)
            base_subtotal = clean_numeric(ws_apu.cell(r, base_group_col + 3).value) or (base_qty * base_price)

            ejec_qty = clean_numeric(ws_apu.cell(r, ejecutada_col).value)
            ejec_unit_qty = clean_numeric(ws_apu.cell(r, unitaria_col).value)
            ejec_price = clean_numeric(ws_apu.cell(r, precio_col).value)
            ejec_subtotal = ejec_qty * ejec_price

            current_apu["resources"].append({
                "insumo": insumo_str.replace(".0", "").strip(),
                "name": nombre,
                "unit": unidad,
                "base_qty": base_qty,
                "base_unit_qty": base_unit_qty,
                "base_price": base_price,
                "base_subtotal": base_subtotal,
                "ejec_qty": ejec_qty,
                "ejec_unit_qty": ejec_unit_qty,
                "ejec_price": ejec_price,
                "ejec_subtotal": ejec_subtotal,
                "excel_faltante_qty": 0.0,
                "excel_faltante_unit_qty": 0.0,
                "excel_faltante_price": 0.0,
                "excel_faltante_subtotal": 0.0
            })

    # Proyecciones y Reutilización
    material_balances = {}
    alerts = []
    reutilization_tips = []

    for apu in apus:
        q_rem = apu["obra_faltante_qty"]

        for res in apu["resources"]:
            ins_code = res["insumo"]
            ins_name = res["name"]
            ins_unit = res["unit"]

            proj_theo_rem_qty = q_rem * res["base_unit_qty"]
            proj_theo_rem_cost = proj_theo_rem_qty * res["base_price"]
            res["proj_theo"] = {
                "qty": res["ejec_qty"] + proj_theo_rem_qty,
                "cost": (res["ejec_qty"] * res["ejec_price"]) + proj_theo_rem_cost,
                "dev_qty": (res["ejec_qty"] + proj_theo_rem_qty) - res["base_qty"],
                "dev_cost": ((res["ejec_qty"] * res["ejec_price"]) + proj_theo_rem_cost) - (res["base_qty"] * res["base_price"])
            }

            u_proj = res["ejec_unit_qty"] if res["ejec_qty"] > 0 else res["base_unit_qty"]
            p_proj = res["ejec_price"] if res["ejec_qty"] > 0 else res["base_price"]

            proj_hist_rem_qty = q_rem * u_proj
            proj_hist_rem_cost = proj_hist_rem_qty * p_proj
            res["proj_hist"] = {
                "qty": res["ejec_qty"] + proj_hist_rem_qty,
                "cost": (res["ejec_qty"] * res["ejec_price"]) + proj_hist_rem_cost,
                "dev_qty": (res["ejec_qty"] + proj_hist_rem_qty) - res["base_qty"],
                "dev_cost": ((res["ejec_qty"] * res["ejec_price"]) + proj_hist_rem_cost) - (res["base_qty"] * res["base_price"])
            }

            dev_qty = res["proj_hist"]["dev_qty"]
            if abs(dev_qty) > 0.01:
                if ins_code not in material_balances:
                    material_balances[ins_code] = {
                        "name": ins_name,
                        "unit": ins_unit,
                        "surpluses": [],
                        "deficits": []
                    }
                if dev_qty < 0:
                    material_balances[ins_code]["surpluses"].append({
                        "apu_code": apu["code"],
                        "apu_name": apu["name"],
                        "qty": abs(dev_qty)
                    })
                else:
                    material_balances[ins_code]["deficits"].append({
                        "apu_code": apu["code"],
                        "apu_name": apu["name"],
                        "qty": dev_qty
                    })

            dev_cost = res["proj_hist"]["dev_cost"]
            base_cost = res["base_qty"] * res["base_price"]
            if dev_cost > 500000:
                pct = (dev_cost / base_cost * 100) if base_cost > 0 else 100
                alerts.append({
                    "type": "danger",
                    "title": "Sobrecosto Crítico en Insumo",
                    "message": f"El insumo '{res['name']}' en el APU '{apu['name']}' proyecta un sobrecosto de {pct:.1f}% ({dev_cost:,.0f} COP) sobre el presupuesto base.",
                    "apu_code": apu["code"],
                    "insumo_code": res["insumo"]
                })

            if res["ejec_qty"] > 0 and res["base_unit_qty"] > 0 and res["ejec_unit_qty"] > res["base_unit_qty"] * 1.15:
                alerts.append({
                    "type": "warning",
                    "title": "Desviación de Rendimiento",
                    "message": f"El rendimiento ejecutado de '{res['name']}' en '{apu['name']}' es un {((res['ejec_unit_qty']/res['base_unit_qty'] - 1)*100):.1f}% superior al presupuesto base.",
                    "apu_code": apu["code"],
                    "insumo_code": res["insumo"]
                })

    for ins_code, balance in material_balances.items():
        surpluses = balance["surpluses"]
        deficits = balance["deficits"]
        if surpluses and deficits:
            for sur in surpluses:
                for defic in deficits:
                    matched_qty = min(sur["qty"], defic["qty"])
                    if matched_qty > 0.01:
                        reutilization_tips.append({
                            "insumo_code": ins_code,
                            "insumo_name": balance["name"],
                            "unit": balance["unit"],
                            "from_apu_code": sur["apu_code"],
                            "from_apu_name": sur["apu_name"],
                            "to_apu_code": defic["apu_code"],
                            "to_apu_name": defic["apu_name"],
                            "qty": matched_qty,
                            "message": f"El material '{balance['name']}' presenta un excedente de {sur['qty']:.1f} {balance['unit']} en '{sur['apu_name']}'. Se recomienda reutilizar {matched_qty:.1f} {balance['unit']} para cubrir el faltante en '{defic['apu_name']}'."
                        })
                        sur["qty"] -= matched_qty
                        defic["qty"] -= matched_qty

    # 6. Codificar archivo en Base64
    out_buffer = io.BytesIO()
    wb.save(out_buffer)
    out_buffer.seek(0)
    excel_b64 = base64.b64encode(out_buffer.getvalue()).decode('utf-8')

    total_base = sum(sum(r["base_qty"] * r["base_price"] for r in a["resources"]) for a in apus)
    total_ejec = sum(sum(r["ejec_qty"] * r["ejec_price"] for r in a["resources"]) for a in apus)
    total_proj_theo = sum(sum(r["proj_theo"]["cost"] for r in a["resources"]) for a in apus)
    total_proj_hist = sum(sum(r["proj_hist"]["cost"] for r in a["resources"]) for a in apus)

    summary = {
        "project_name": filename.replace("Copia de V1 SAO Costos por niveles_", "").split(".")[0].replace("_", " ").strip(),
        "total_base": total_base,
        "total_ejec": total_ejec,
        "total_proj_theo": total_proj_theo,
        "total_proj_hist": total_proj_hist,
        "dev_proj_theo": total_proj_theo - total_base,
        "dev_proj_hist": total_proj_hist - total_base,
        "total_items": len(apus),
        "total_alerts": len(alerts),
        "total_reutilizations": len(reutilization_tips)
    }

    return {
        "summary": summary,
        "alerts": alerts,
        "reutilizaciones": reutilization_tips,
        "details": apus,
        "excel_b64": excel_b64
    }
